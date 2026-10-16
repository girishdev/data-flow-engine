
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import load_workbook


REQUIRED_HEADERS = [
    "signal_name",
    "application",
    "signal_direction",
    "factor",
    "safety",
    "status",
    "description",
]


class ExcelValidationError(Exception):
    pass


def extract_excel_data(file_content: bytes) -> dict:
    workbook = load_workbook(
        BytesIO(file_content),
        read_only=True,
        data_only=True
    )

    try:
        if len(workbook.worksheets) != 1:
            raise ExcelValidationError(
                "Excel must contain exactly one worksheet"
            )

        sheet = workbook.worksheets[0]

        rows = sheet.iter_rows(values_only=True)

        # First row contains the column headers
        headers = next(rows, None)

        if headers is None:
            raise ExcelValidationError("Excel sheet is empty")

        headers = [
            str(value).strip() if value is not None else ""
            for value in headers
        ]

        if headers != REQUIRED_HEADERS:
            raise ExcelValidationError(
                "Invalid Excel headers. Expected: "
                + ", ".join(REQUIRED_HEADERS)
            )

        records = []

        for row in rows:
            # Ignore completely empty rows
            if all(value is None for value in row):
                continue

            record = {}

            for header, value in zip(headers, row):
                if isinstance(value, (datetime, date, time)):
                    value = value.isoformat()

                if not isinstance(
                    value, (str, int, float, bool, type(None))
                ):
                    raise ExcelValidationError(
                        f"Unsupported cell value in column: {header}"
                    )

                record[header] = value

            records.append(record)

        return {
            "sheet_name": sheet.title,
            "total_records": len(records),
            "data": records
        }

    finally:
        workbook.close()
