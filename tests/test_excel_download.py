
from io import BytesIO
from types import SimpleNamespace

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.database.connection import get_db
from app.routes import files
from app.services.excel_export_service import generate_excel


SAMPLE_DATA = {
    "sheet_name": "Signals",
    "total_records": 1,
    "data": [
        {
            "signal_name": "EngineSpeed",
            "application": "Engine",
            "signal_direction": "Input",
            "factor": 1.5,
            "safety": True,
            "status": "Active",
            "description": "Engine signal"
        }
    ]
}


def test_generate_excel():
    output = generate_excel(SAMPLE_DATA)

    workbook = load_workbook(output, data_only=True)
    sheet = workbook["Signals"]

    assert sheet["A1"].value == "signal_name"
    assert sheet["A2"].value == "EngineSpeed"
    assert sheet["D2"].value == 1.5
    assert sheet["E2"].value is True
    assert sheet["F2"].value == "Active"

    workbook.close()


def test_download_success(monkeypatch):
    request_id = "DFE-TEST123"

    def fake_lookup(db, request_id):
        return SimpleNamespace(
            request_id=request_id,
            data=SAMPLE_DATA
        )

    monkeypatch.setattr(
        files,
        "get_automation_by_request_id",
        fake_lookup
    )

    app.dependency_overrides[get_db] = lambda: None

    try:
        with TestClient(app) as client:
            response = client.get(
                f"/api/files/{request_id}/download"
            )

        assert response.status_code == 200

        assert "attachment" in response.headers[
            "content-disposition"
        ]

        workbook = load_workbook(
            BytesIO(response.content)
        )
        assert workbook["Signals"]["A2"].value == "EngineSpeed"
        workbook.close()

    finally:
        app.dependency_overrides.clear()


def test_invalid_request_id(monkeypatch):
    def fake_lookup(db, request_id):
        return None

    monkeypatch.setattr(
        files,
        "get_automation_by_request_id",
        fake_lookup
    )

    app.dependency_overrides[get_db] = lambda: None

    try:
        with TestClient(app) as client:
            response = client.get(
                "/api/files/INVALID-ID/download"
            )

        assert response.status_code == 404
        assert response.json()["detail"] == "Request ID not found"

    finally:
        app.dependency_overrides.clear()
